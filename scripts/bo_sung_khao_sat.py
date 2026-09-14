from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"C:\Users\ndhng\Downloads\YOLO_Pose_and_ST_GCN")
SRC = ROOT / "docs" / "KhaoSatHienTrangPhatHienTeNga.docx"
OUT = ROOT / "docs" / "KhaoSatHienTrangPhatHienTeNga_DaBoSung.docx"
ASSETS = ROOT / ".tmp_doc_assets"
ASSETS.mkdir(exist_ok=True)

def make_diagram(name, labels):
    im = Image.new("RGB", (1600, 500), "white")
    d = ImageDraw.Draw(im); f = ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", 24)
    n, w, h, y = len(labels), 270, 110, 160
    gap = (1600 - 120 - n*w) // max(n-1, 1)
    for i, label in enumerate(labels):
        x = 60 + i*(w+gap)
        d.rounded_rectangle((x,y,x+w,y+h), 16, fill="#EAF2F8", outline="#1F4E78", width=4)
        lines = label.split("\n"); yy = y + 22
        for line in lines:
            bb=d.textbbox((0,0),line,font=f); d.text((x+(w-bb[2]+bb[0])//2,yy),line,font=f,fill="#17365D"); yy += 34
        if i < n-1:
            x1=x+w+8; x2=60+(i+1)*(w+gap)-8; mid=y+h//2
            d.line((x1,mid,x2,mid),fill="#1F4E78",width=5); d.polygon([(x2,mid),(x2-18,mid-11),(x2-18,mid+11)],fill="#1F4E78")
    im.save(ASSETS / name)

make_diagram("tong_quan.png", ["Camera\ntrong phòng", "Máy xử lý\ntại chỗ", "Cảnh báo\ntại chỗ", "Người chăm sóc\nnhận thông báo"])
make_diagram("canh_bao.png", ["Nghi ngờ\nté ngã", "Hỏi/xác nhận\ntại chỗ", "Gửi người\nliên hệ 1", "Leo thang\nngười tiếp theo"])
make_diagram("han_che.png", ["Mất kết nối /\nánh sáng yếu", "Báo trạng thái\nhạn chế", "Người phụ trách\nkiểm tra", "Khôi phục &\nghi nhật ký"])

doc = Document(SRC)
paras = {p.text.strip(): p for p in doc.paragraphs if p.text.strip()}

def find_prefix(prefix):
    for text, p in paras.items():
        if text.startswith(prefix):
            return p
    raise KeyError(prefix)

def before_heading(prefix):
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith(prefix):
            return doc.paragraphs[i - 1]
    raise KeyError(prefix)

def style(p, bold=False, center=False):
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(6)
    for r in p.runs:
        r.font.name = "Times New Roman"; r.font.size = Pt(13); r.bold = bold
        r._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        r._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")

def paragraph(text, bold=False, center=False):
    p=doc.add_paragraph(text); style(p,bold,center); return p

def after(anchor, items):
    a=anchor._p
    for item in items:
        el=item._p if hasattr(item,"_p") else item._tbl
        a.addnext(el); a=el

def shade(cell):
    x=OxmlElement("w:shd"); x.set(qn("w:fill"),"D9EAF7"); cell._tc.get_or_add_tcPr().append(x)

def table(headers, rows):
    t=doc.add_table(rows=1, cols=len(headers)); t.style="Table Grid"; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for j,h in enumerate(headers):
        c=t.rows[0].cells[j]; c.text=h; shade(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs: r.bold=True; r.font.name="Times New Roman"; r.font.size=Pt(10.5)
    for row in rows:
        cs=t.add_row().cells
        for j,val in enumerate(row):
            cs[j].text=val; cs[j].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cs[j].paragraphs:
                p.paragraph_format.space_after=Pt(2)
                for r in p.runs: r.font.name="Times New Roman"; r.font.size=Pt(10)
    return t

a4=paras["Trong các hướng trên, camera kết hợp phân tích chuyển động phù hợp nhất với mục tiêu đề tài vì có thể tận dụng camera sẵn có trong nhà và không yêu cầu người cao tuổi thực hiện thao tác. Tuy nhiên, cách tiếp cận này chỉ có ý nghĩa khi việc xử lý hình ảnh được thực hiện cục bộ, không lưu trữ video thô không cần thiết và có cơ chế kiểm soát cảnh báo sai."]
p1=paragraph("A.4.1. Một số hệ thống và sản phẩm tiêu biểu trên thị trường", True)
p2=paragraph("Để tránh chỉ dừng ở việc phân loại giải pháp theo lý thuyết, nhóm khảo sát thêm một số sản phẩm đang được công bố và sử dụng. Apple Watch Fall Detection đại diện cho hướng thiết bị đeo: đồng hồ rung, phát âm báo và hiển thị cảnh báo khi nhận biết cú ngã mạnh; nếu người dùng bất động khoảng một phút, thiết bị có thể gọi trợ giúp và báo cho liên hệ khẩn cấp. Vayyar Care đại diện cho hướng cảm biến môi trường không dùng camera; hệ thống sử dụng cảm biến radar để theo dõi không chạm, nhấn mạnh việc giảm lo ngại về hình ảnh riêng tư. Nobi Smart Lights là hướng tích hợp công nghệ chăm sóc vào đèn chiếu sáng trong phòng, kết hợp phòng ngừa ngã, phát hiện ngã và gửi cảnh báo cho người được chỉ định.")
p3=paragraph("Các ví dụ này cho thấy một hệ thống hữu ích không chỉ dừng ở thuật toán phát hiện. Điểm đáng kế thừa là cơ chế cảnh báo rõ ràng, có thời gian để người dùng phản hồi và chuyển thông tin đến đúng người cần hỗ trợ. Tuy nhiên, thiết bị đeo vẫn phụ thuộc vào việc đeo và sạc; cảm biến môi trường hoặc đèn thông minh thường đòi hỏi đầu tư phần cứng chuyên dụng; còn các dịch vụ đóng gói sẵn có thể hạn chế khả năng tùy biến cho từng phòng và từng quy trình chăm sóc. Đề tài chọn hướng tận dụng camera IP phổ thông, xử lý tại nơi lắp đặt và thiết kế cơ chế xác nhận/leo thang có thể cấu hình để khắc phục các điểm này.")
t1=table(["Hệ thống", "Điểm đáng chú ý", "Giới hạn", "Kế thừa/khắc phục"], [["Apple Watch Fall Detection", "Cảnh báo tại cổ tay; tự liên hệ trợ giúp khi bất động kéo dài.", "Phụ thuộc việc đeo, sạc và thiết bị còn hoạt động.", "Kế thừa xác nhận trước khi leo thang; không buộc người cao tuổi mang thiết bị."], ["Vayyar Care", "Cảm biến radar không dùng camera, hướng đến theo dõi không chạm và riêng tư.", "Cần cảm biến chuyên dụng và khảo sát vị trí lắp đặt.", "Kế thừa yêu cầu riêng tư; dùng xử lý cục bộ và chỉ giữ dữ liệu cần thiết."], ["Nobi Smart Lights", "Tích hợp cảnh báo chăm sóc vào thiết bị trong phòng; hỗ trợ thông báo cho người được chỉ định.", "Cần thay thế/bổ sung thiết bị chiếu sáng chuyên dụng.", "Kế thừa thông báo đúng người; tận dụng camera có sẵn để giảm chi phí ban đầu."]])
after(a4,[p1,p2,p3,t1])

a5=paras["A.5. Khoảng trống và định hướng khắc phục"]
p4=paragraph("Từ khảo sát trên, điểm sáng tạo của hệ thống không được hiểu là thay thế hoàn toàn vai trò chăm sóc bằng máy móc. Hệ thống được đề xuất kết hợp bốn ý tưởng: phát hiện chủ động từ diễn biến tư thế thay vì chỉ báo chuyển động; xử lý cục bộ để hạn chế đưa video thô ra ngoài; xác nhận tại chỗ trước khi làm phiền người chăm sóc; và leo thang cảnh báo theo thứ tự liên hệ khi chưa có người tiếp nhận. Ngoài ra, hệ thống chủ động công khai trạng thái “giám sát hạn chế” khi camera, ánh sáng hoặc kết nối không bảo đảm, qua đó tránh tạo cảm giác an toàn giả.")
after(a5,[p4])

b1=before_heading("B.2.")
img1=doc.add_paragraph(); img1.alignment=WD_ALIGN_PARAGRAPH.CENTER; img1.add_run().add_picture(str(ASSETS/"tong_quan.png"),width=Inches(6.4))
cap1=paragraph("Hình 1. Sơ đồ khối tổng quan hoạt động của hệ thống", center=True)
after(b1,[img1,cap1])

b42=before_heading("B.4.3.")
img2=doc.add_paragraph(); img2.alignment=WD_ALIGN_PARAGRAPH.CENTER; img2.add_run().add_picture(str(ASSETS/"canh_bao.png"),width=Inches(6.4))
cap2=paragraph("Hình 2. Chuỗi xác nhận và leo thang cảnh báo", center=True)
p5=paragraph("Các tình huống cần được quy định rõ như sau: nếu người cao tuổi phản hồi hoặc đứng dậy, sự kiện được ghi nhận là đã tự xác nhận; nếu người liên hệ thứ nhất không xác nhận trong thời hạn, hệ thống chuyển sang người tiếp theo; nếu tất cả người liên hệ chưa phản hồi, hệ thống duy trì cảnh báo ở mức cao và thực hiện hướng dẫn đã được cơ sở/hộ gia đình thống nhất trước, chẳng hạn gọi điện trực tiếp hoặc liên hệ số hỗ trợ khẩn cấp. Nếu có nhiều người trong khung hình, hệ thống không tự gán sự cố cho một người khi độ tin cậy không đủ mà chuyển sang trạng thái cần kiểm tra. Việc gọi cơ quan y tế khẩn cấp không do hệ thống tự quyết định, trừ khi chủ sở hữu đã thiết lập và được phép theo quy định áp dụng.")
after(b42,[img2,cap2,p5])

b43=before_heading("B.5.")
p6=paragraph("B.4.4. Xử lý tình trạng giám sát hạn chế", True)
p7=paragraph("Khi mất điện, mất kết nối camera, Internet gián đoạn, hình ảnh quá tối hoặc vùng quan sát bị che khuất, hệ thống phải tách hai việc: chức năng phát hiện cục bộ chỉ tiếp tục khi camera và máy xử lý vẫn hoạt động; còn việc gửi thông báo ra ngoài có thể chờ mạng được khôi phục. Trong mọi trường hợp, màn hình quản trị và người phụ trách phải thấy trạng thái giám sát hạn chế. Khi sự cố được khắc phục, hệ thống ghi nhận thời điểm bắt đầu/kết thúc để người dùng biết khoảng thời gian nào không thể tin cậy hoàn toàn vào cơ chế tự động.")
img3=doc.add_paragraph(); img3.alignment=WD_ALIGN_PARAGRAPH.CENTER; img3.add_run().add_picture(str(ASSETS/"han_che.png"),width=Inches(6.4))
cap3=paragraph("Hình 3. Quy trình xử lý trạng thái giám sát hạn chế", center=True)
after(b43,[p6,p7,img3,cap3])

ref=paras["[8] Auvinet, E. và cộng sự. Multiple cameras fall dataset. Université de Montréal."]
r9=paragraph("[9] Apple Support. Use Fall Detection with Apple Watch. https://support.apple.com/en-la/108896. Truy cập ngày 26/07/2026.")
r10=paragraph("[10] Vayyar. Vayyar Care: Fall Detection. https://vayyar.com/care-docs/b2c/. Truy cập ngày 26/07/2026.")
r11=paragraph("[11] Nobi. Nobi Smart Lights – Fall Prevention and Fall Detection. https://www.nobi.life/en. Truy cập ngày 26/07/2026.")
after(ref,[r9,r10,r11])

doc.save(OUT)
print(OUT)
