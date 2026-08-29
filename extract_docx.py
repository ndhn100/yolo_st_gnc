from docx import Document
import sys

sys.stdout.reconfigure(encoding="utf-8")

for filename in sys.argv[1:]:
    print(f"\n--- {filename} ---")
    document = Document(filename)
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            print(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            print("TABLE: " + " | ".join(cell.text.replace("\n", " / ") for cell in row.cells))
